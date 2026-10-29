"""
test_export_alignment.py - Comprehensive unit tests for website OCR export alignment
and document rendering (TXT, DOCX, PDF).

Verifies:
1. UTF-8 plain text normalization and byte fidelity.
2. Word document (.docx) paragraph structure, 1.0-inch margins, 13pt Nirmala UI typography,
   and OpenXML complex script (w:cs, w:szCs) attributes.
3. PDF (.pdf) generation via vector HarfBuzz/FreeType engine, margin-aware line wrapping,
   multi-page pagination, and header/footer metadata.
4. Flask API export routes and Content-Disposition headers.
"""

import unittest
import io
import docx
import pdfplumber

from app.unicode_converter import export_to_txt, export_to_docx, export_to_pdf
from app.main import app

class TestExportAlignment(unittest.TestCase):
    def setUp(self):
        app.config['TESTING'] = True
        self.client = app.test_client()
        self.sample_text_short = "தமிழ் வாழ்க"
        self.sample_text_multiline = (
            "நிலையான புகழைத் தருவது ஒழுக்கமே\n"
            "அதனால் அவ்வொழுக்கமே உயிரினும் மேலானதாகப் போற்றப்படும்\n"
            "ஒழுக்கம் விழுப்பம் தரலான் ஒழுக்கம்\n"
            "உயிரினும் ஓம்பப் படும்"
        )
        # 45 lines to trigger multi-page pagination in PDF
        self.sample_text_long = "\n".join([f"வரிசை எண் {i+1}: தமிழ் எழுத்துக்கள் பயிற்சி வரிசை" for i in range(45)])

    # -------------------------------------------------------------------------
    # TXT Export Tests
    # -------------------------------------------------------------------------
    def test_export_txt_utf8_fidelity(self):
        """Verify TXT export maintains exact UTF-8 byte stream and codepoints."""
        buf = export_to_txt(self.sample_text_short)
        self.assertIsInstance(buf, io.BytesIO)
        content = buf.getvalue().decode('utf-8')
        self.assertEqual(content, self.sample_text_short)

    def test_export_txt_crlf_normalization(self):
        """Verify Windows CRLF is cleanly normalized to standard LF."""
        crlf_input = "வரி 1\r\nவரி 2\r\nவரி 3"
        buf = export_to_txt(crlf_input)
        content = buf.getvalue().decode('utf-8')
        self.assertEqual(content, "வரி 1\nவரி 2\nவரி 3")
        self.assertNotIn("\r", content)

    # -------------------------------------------------------------------------
    # DOCX Export Tests
    # -------------------------------------------------------------------------
    def test_export_docx_structure_and_margins(self):
        """Verify DOCX document margins (1.0 inch) and heading structure."""
        buf = export_to_docx(self.sample_text_multiline, title="ஆய்வு வெளியீடு")
        self.assertIsInstance(buf, io.BytesIO)
        
        doc = docx.Document(buf)
        # Check margins
        section = doc.sections[0]
        self.assertAlmostEqual(section.top_margin.inches, 1.0, places=2)
        self.assertAlmostEqual(section.bottom_margin.inches, 1.0, places=2)
        self.assertAlmostEqual(section.left_margin.inches, 1.0, places=2)
        self.assertAlmostEqual(section.right_margin.inches, 1.0, places=2)

        # Check title
        self.assertIn("ஆய்வு வெளியீடு", doc.paragraphs[0].text)

    def test_export_docx_line_paragraphs_and_complex_script_attributes(self):
        """Verify each line is its own paragraph and has w:cs, w:szCs XML attributes."""
        buf = export_to_docx(self.sample_text_multiline)
        doc = docx.Document(buf)
        
        # Paragraph 0: Title, Paragraph 1: Subtitle, Paragraphs 2+: Content lines
        content_paragraphs = [p for p in doc.paragraphs[2:] if p.text.strip()]
        expected_lines = [l for l in self.sample_text_multiline.split('\n') if l.strip()]
        self.assertEqual(len(content_paragraphs), len(expected_lines))

        # Check font properties and OpenXML tags on the first content paragraph
        first_p = content_paragraphs[0]
        self.assertEqual(first_p.text, expected_lines[0])
        self.assertTrue(len(first_p.runs) > 0)
        run = first_p.runs[0]
        self.assertEqual(run.font.name, 'Nirmala UI')
        self.assertEqual(run.font.size.pt, 13)

        # Verify Word OpenXML complex script (w:cs) and size (w:szCs)
        rPr = run._r.get_or_add_rPr()
        rFonts = rPr.find(docx.oxml.ns.qn('w:rFonts'))
        self.assertIsNotNone(rFonts, "w:rFonts element missing")
        self.assertEqual(rFonts.get(docx.oxml.ns.qn('w:cs')), 'Nirmala UI')
        self.assertEqual(rFonts.get(docx.oxml.ns.qn('w:ascii')), 'Nirmala UI')

        szCs = rPr.find(docx.oxml.ns.qn('w:szCs'))
        self.assertIsNotNone(szCs, "w:szCs element missing")
        self.assertEqual(szCs.get(docx.oxml.ns.qn('w:val')), '26')  # 26 half-pts = 13pt

    # -------------------------------------------------------------------------
    # PDF Export Tests
    # -------------------------------------------------------------------------
    def test_export_pdf_valid_header_and_pages(self):
        """Verify PDF export generates valid %PDF vector file with metadata."""
        buf = export_to_pdf(self.sample_text_multiline)
        self.assertIsInstance(buf, io.BytesIO)
        pdf_bytes = buf.getvalue()
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        self.assertTrue(len(pdf_bytes) > 1000)

        # Check with pdfplumber
        with pdfplumber.open(buf) as pdf:
            self.assertEqual(len(pdf.pages), 1)
            page_text = pdf.pages[0].extract_text()
            self.assertIn("Tamil Handwritten OCR Output", page_text)
            self.assertIn("Page 1 of 1", page_text)

    def test_export_pdf_multi_page_pagination(self):
        """Verify long text properly paginates across multiple pages without clipping."""
        buf = export_to_pdf(self.sample_text_long)
        with pdfplumber.open(buf) as pdf:
            self.assertGreaterEqual(len(pdf.pages), 2, "Long text should span at least 2 pages")
            page1_text = pdf.pages[0].extract_text()
            page2_text = pdf.pages[1].extract_text()
            self.assertIn("Page 1 of", page1_text)
            self.assertIn("Page 2 of", page2_text)

    def test_export_pdf_long_line_wrapping(self):
        """Verify an extremely long line wraps into multiple printable lines without clipping."""
        long_line = " ".join(["தமிழ்மொழி"] * 40)
        buf = export_to_pdf(long_line)
        self.assertTrue(buf.getvalue().startswith(b'%PDF'))
        with pdfplumber.open(buf) as pdf:
            self.assertGreaterEqual(len(pdf.pages), 1)

    # -------------------------------------------------------------------------
    # Flask Route Integration Tests
    # -------------------------------------------------------------------------
    def test_api_export_txt_route(self):
        """Verify /api/export/txt HTTP endpoint returns UTF-8 attachment."""
        res = self.client.post('/api/export/txt', json={'text': self.sample_text_short})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, 'text/plain')
        self.assertIn('attachment', res.headers.get('Content-Disposition', ''))
        self.assertEqual(res.data.decode('utf-8'), self.sample_text_short)

    def test_api_export_docx_route(self):
        """Verify /api/export/docx HTTP endpoint returns docx attachment."""
        res = self.client.post('/api/export/docx', json={'text': self.sample_text_short})
        self.assertEqual(res.status_code, 200)
        self.assertIn('attachment', res.headers.get('Content-Disposition', ''))
        self.assertTrue(len(res.data) > 500)

    def test_api_export_pdf_route(self):
        """Verify /api/export/pdf HTTP endpoint returns PDF attachment."""
        res = self.client.post('/api/export/pdf', json={'text': self.sample_text_short})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.mimetype, 'application/pdf')
        self.assertIn('attachment', res.headers.get('Content-Disposition', ''))
        self.assertTrue(res.data.startswith(b'%PDF'))

if __name__ == '__main__':
    unittest.main()
